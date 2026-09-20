"""Boutframe exports must include subjects from combined (merged) projects."""

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from abel.models.schemas import (
    CandidateWindow,
    ImportManifest,
    LinkedSession,
    ReviewDecision,
    ReviewDecisionType,
    VideoAsset,
)
from abel.services.export_service import ExportService
from abel.services.import_service import ImportService


def _write_project(root: Path, subject: str, session_id: str) -> None:
    manifest = ImportManifest(
        videos=[VideoAsset(asset_id="v1", source_path=f"{subject}.avi", subject_id=subject)],
        linked_sessions=[
            LinkedSession(
                session_id=session_id,
                video_asset_id="v1",
                pose_asset_id="p1",
                subject_id=subject,
            )
        ],
    )
    ImportService().save_manifest(root, manifest)


def _write_external_bouts(root: Path, session_id: str, bouts: list[tuple[int, int]]) -> None:
    (root / "config").mkdir(parents=True, exist_ok=True)
    (root / "config" / "behavior_definitions.yaml").write_text(
        "behaviors:\n  - behavior_id: rear\n    name: rear\n",
        encoding="utf-8",
    )
    bouts_dir = root / "derived" / "behavior_bouts"
    bouts_dir.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        [
            {"session_id": session_id, "start_frame": s, "end_frame": e}
            for s, e in bouts
        ]
    ).to_parquet(bouts_dir / "rear_bouts.parquet")


def _host_project(tmp_path: Path) -> tuple[Path, list[CandidateWindow], list[ReviewDecision]]:
    host = tmp_path / "host"
    (host / "derived" / "review_tables").mkdir(parents=True)
    _write_project(host, subject="HostMouse", session_id="s1")

    candidates = [
        CandidateWindow(
            window_id="c1", session_id="s1", start_frame=10, end_frame=20, behavior_id="rear"
        )
    ]
    decisions = [
        ReviewDecision(
            decision_id="d1",
            clip_id="c1",
            reviewer="r",
            old_status="unscored",
            new_status="reviewed",
            decision=ReviewDecisionType.ACCEPT,
            adjusted_start_frame=11,
            adjusted_end_frame=19,
        )
    ]
    return host, candidates, decisions


def _register_merged(host: Path, external: Path, tag: str) -> None:
    (host / "config").mkdir(parents=True, exist_ok=True)
    (host / "config" / "merged_projects.json").write_text(
        json.dumps([{"project_root": str(external), "tag": tag, "group_override": ""}]),
        encoding="utf-8",
    )


def test_boutframes_export_includes_merged_project_subjects(tmp_path: Path) -> None:
    host, candidates, decisions = _host_project(tmp_path)

    external = tmp_path / "external"
    (external / "derived" / "review_tables").mkdir(parents=True)
    _write_project(external, subject="ExtMouse", session_id="e1")
    _write_external_bouts(external, session_id="e1", bouts=[(100, 120), (200, 230)])
    _register_merged(host, external, tag="StudyB")

    service = ExportService()
    service.set_project(host)
    out = service.export_boutframes_xlsx(candidates, decisions, filename="boutframes.xlsx")

    assert out.success
    assert out.n_merged_projects == 1

    wb = load_workbook(out.output_path)
    sheets = [name for name in wb.sheetnames if name != "_bout_counts"]
    assert "HostMouse" in sheets
    # "StudyB/ExtMouse": the forward slash is not legal in an Excel sheet title.
    assert "StudyB_ExtMouse" in sheets

    ext_rows = list(wb["StudyB_ExtMouse"].iter_rows(values_only=True))
    assert ext_rows[0] == ("rear",)
    assert [row[0] for row in ext_rows[1:]] == [100, 200]

    counts = {row[0]: row[1] for row in wb["_bout_counts"].iter_rows(min_row=2, values_only=True)}
    assert counts["StudyB/ExtMouse"] == 2
    assert counts["HostMouse"] == 1


def test_boutframes_export_can_exclude_merged_projects(tmp_path: Path) -> None:
    host, candidates, decisions = _host_project(tmp_path)

    external = tmp_path / "external"
    (external / "derived" / "review_tables").mkdir(parents=True)
    _write_project(external, subject="ExtMouse", session_id="e1")
    _write_external_bouts(external, session_id="e1", bouts=[(100, 120)])
    _register_merged(host, external, tag="StudyB")

    service = ExportService()
    service.set_project(host)
    out = service.export_boutframes_xlsx(
        candidates, decisions, filename="boutframes.xlsx", include_merged_projects=False
    )

    assert out.success
    assert out.n_merged_projects == 0
    wb = load_workbook(out.output_path)
    assert [name for name in wb.sheetnames if name != "_bout_counts"] == ["HostMouse"]


def test_boutframes_export_survives_missing_merged_project(tmp_path: Path) -> None:
    host, candidates, decisions = _host_project(tmp_path)
    _register_merged(host, tmp_path / "gone", tag="StudyB")

    service = ExportService()
    service.set_project(host)
    out = service.export_boutframes_xlsx(candidates, decisions, filename="boutframes.xlsx")

    assert out.success
    assert out.n_merged_projects == 0
    wb = load_workbook(out.output_path)
    assert [name for name in wb.sheetnames if name != "_bout_counts"] == ["HostMouse"]
